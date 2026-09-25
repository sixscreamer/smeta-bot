import os
import io
import re
import html
import logging
from decimal import Decimal, InvalidOperation

import asyncpg
from telegram import (
    Update, InlineKeyboardButton, InlineKeyboardMarkup,
)
from telegram.constants import ParseMode
from telegram.ext import (
    Application, CommandHandler, CallbackQueryHandler,
    MessageHandler, filters, ContextTypes,
)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer,
)

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(name)s | %(message)s",
)
log = logging.getLogger("smeta")

BOT_TOKEN = os.environ["BOT_TOKEN"]
DATABASE_URL = os.environ["DATABASE_URL"]

pool: asyncpg.Pool | None = None

CATEGORIES = [
    "🎬 Аренда",
    "👕 Одежда",
    "💄 Грим",
    "🎨 Реквизит",
    "🚕 Транспорт",
    "🍔 Кейтеринг",
    "👥 Команда",
    "📦 Другое",
]

# Роли заранее — пригодятся, когда будем ограничивать права
ROLE_OWNER = "owner"
ROLE_MEMBER = "member"
ROLE_VIEWER = "viewer"


# ---------- DB ----------

SCHEMA = """
CREATE TABLE IF NOT EXISTS users (
    user_id     BIGINT PRIMARY KEY,
    username    TEXT,
    first_name  TEXT,
    last_name   TEXT,
    updated_at  TIMESTAMPTZ NOT NULL DEFAULT NOW()
);

CREATE TABLE IF NOT EXISTS projects (
    id          SERIAL PRIMARY KEY,
    name        TEXT NOT NULL,
    budget      NUMERIC(14,2) NOT NULL DEFAULT 0,
    creator_id  BIGINT NOT NULL,
    created_at  TIMESTAMPTZ NOT NULL DEFAULT NOW()
);

CREATE TABLE IF NOT EXISTS project_members (
    project_id  INTEGER NOT NULL REFERENCES projects(id) ON DELETE CASCADE,
    user_id     BIGINT NOT NULL,
    role        TEXT NOT NULL DEFAULT 'member',
    added_at    TIMESTAMPTZ NOT NULL DEFAULT NOW(),
    PRIMARY KEY (project_id, user_id)
);

CREATE TABLE IF NOT EXISTS expenses (
    id              SERIAL PRIMARY KEY,
    project_id      INTEGER NOT NULL REFERENCES projects(id) ON DELETE CASCADE,
    category        TEXT NOT NULL DEFAULT '📦 Другое',
    name            TEXT NOT NULL,
    qty             NUMERIC(14,3) NOT NULL DEFAULT 1,
    price           NUMERIC(14,2) NOT NULL DEFAULT 0,
    comment         TEXT,
    author_id       BIGINT,
    author_username TEXT,
    created_at      TIMESTAMPTZ NOT NULL DEFAULT NOW()
);

CREATE INDEX IF NOT EXISTS idx_expenses_project ON expenses(project_id);
CREATE INDEX IF NOT EXISTS idx_members_user     ON project_members(user_id);
"""


async def init_db():
    global pool
    pool = await asyncpg.create_pool(DATABASE_URL, min_size=1, max_size=5)
    async with pool.acquire() as c:
        await c.execute(SCHEMA)
    log.info("DB schema ready")


async def upsert_user(u):
    if u is None:
        return
    async with pool.acquire() as c:
        await c.execute(
            """
            INSERT INTO users(user_id, username, first_name, last_name, updated_at)
            VALUES($1,$2,$3,$4,NOW())
            ON CONFLICT (user_id) DO UPDATE
            SET username   = EXCLUDED.username,
                first_name = EXCLUDED.first_name,
                last_name  = EXCLUDED.last_name,
                updated_at = NOW()
            """,
            u.id, u.username, u.first_name, u.last_name,
        )


# ---------- helpers ----------

def money(v) -> str:
    if v is None:
        v = 0
    d = Decimal(str(v))
    sign = "-" if d < 0 else ""
    d = abs(d).quantize(Decimal("0.01"))
    s = f"{d:,.2f}".replace(",", " ").replace(".", ",")
    return f"{sign}{s} ₽"


def fmt_qty(q) -> str:
    d = Decimal(str(q)).normalize()
    if d == d.to_integral_value():
        return str(int(d))
    return str(d).replace(".", ",")


def esc(s) -> str:
    return html.escape(str(s)) if s is not None else ""


def parse_money(text):
    if not text:
        return None
    t = text.strip().lower()
    for token in ("₽", "руб.", "руб", "р.", "rub", "rur"):
        t = t.replace(token, "")
    t = t.replace(" ", "").replace("\u00a0", "")
    if "," in t and "." in t:
        if t.rfind(",") > t.rfind("."):
            t = t.replace(".", "").replace(",", ".")
        else:
            t = t.replace(",", "")
    else:
        t = t.replace(",", ".")
    try:
        d = Decimal(t)
        if d < 0:
            return None
        return d.quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def parse_number(text):
    if not text:
        return None
    t = text.strip().replace(",", ".").replace(" ", "")
    try:
        d = Decimal(t)
        if d < 0:
            return None
        return d
    except InvalidOperation:
        return None


MEASURE_WORDS = {
    "шт", "штук", "штука", "штуки",
    "банок", "банка", "банки",
    "кг", "г", "гр", "л", "мл",
    "м", "см", "мм",
    "упаковок", "упаковка", "упаковки",
    "пачек", "пачка", "пачки",
    "бутылок", "бутылка", "бутылки",
    "коробок", "коробка", "коробки",
    "рулон", "рулона", "рулонов",
}


def _num(s: str) -> Decimal:
    return Decimal(s.replace(",", "."))


def _clean_name(s: str) -> str:
    parts = [p for p in s.strip().split() if p]
    while parts and parts[0].lower() in MEASURE_WORDS:
        parts.pop(0)
    name = " ".join(parts) if parts else s.strip()
    if name:
        name = name[0].upper() + name[1:]
    return name


def parse_quick(text: str):
    """Возвращает (name, qty, price) или None. Если неоднозначно — None."""
    if not text:
        return None
    t = text.strip()
    low = t.lower()

    # 3 банки краски по 850
    m = re.match(
        r"^(\d+(?:[.,]\d+)?)\s+(.+?)\s+по\s+(\d+(?:[.,]\d+)?)\s*"
        r"(?:₽|руб\.?|р\.?)?$",
        low,
    )
    if m:
        return _clean_name(m.group(2)), _num(m.group(1)), _num(m.group(3))

    # 3 банки краски 850
    m = re.match(
        r"^(\d+(?:[.,]\d+)?)\s+(.+?)\s+(\d+(?:[.,]\d+)?)\s*"
        r"(?:₽|руб\.?|р\.?)?$",
        low,
    )
    if m:
        return _clean_name(m.group(2)), _num(m.group(1)), _num(m.group(3))

    # краска 3 x 850
    m = re.match(
        r"^(.+?)\s+(\d+(?:[.,]\d+)?)\s*[x×*]\s*(\d+(?:[.,]\d+)?)\s*"
        r"(?:₽|руб\.?|р\.?)?$",
        low,
    )
    if m:
        return _clean_name(m.group(1)), _num(m.group(2)), _num(m.group(3))

    # такси 1200
    m = re.match(
        r"^([^\d].*?)\s+(\d+(?:[.,]\d+)?)\s*(?:₽|руб\.?|р\.?)?$",
        low,
    )
    if m:
        return _clean_name(m.group(1)), Decimal(1), _num(m.group(2))

    return None


# ---------- keyboards ----------

def kb_main():
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("📁 Общие проекты", callback_data="projects")],
        [InlineKeyboardButton("➕ Новый проект", callback_data="project:new")],
    ])


def kb_back(cb="projects"):
    return InlineKeyboardMarkup([[InlineKeyboardButton("◀️ Назад", callback_data=cb)]])


def kb_project(pid):
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("➕ Добавить расход", callback_data=f"exp:add:{pid}")],
        [
            InlineKeyboardButton("📋 Расходы", callback_data=f"exp:list:{pid}"),
            InlineKeyboardButton("📊 Excel",  callback_data=f"excel:{pid}"),
            InlineKeyboardButton("📄 PDF",    callback_data=f"pdf:{pid}"),
        ],
        [
            InlineKeyboardButton("⚙️ Настройки",      callback_data=f"project:settings:{pid}"),
            InlineKeyboardButton("🗑 Удалить проект",  callback_data=f"project:delete:{pid}"),
        ],
        [InlineKeyboardButton("◀️ К списку", callback_data="projects")],
    ])


def kb_categories(pid):
    rows = [[InlineKeyboardButton(c, callback_data=f"cat:{pid}:{i}")] for i, c in enumerate(CATEGORIES)]
    rows.append([InlineKeyboardButton("❌ Отмена", callback_data=f"exp:cancel:{pid}")])
    return InlineKeyboardMarkup(rows)


def kb_comment(pid):
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("⏭ Пропустить", callback_data=f"exp:skip_comment:{pid}")],
        [InlineKeyboardButton("❌ Отмена",     callback_data=f"exp:cancel:{pid}")],
    ])


def kb_confirm(pid):
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("✅ Сохранить", callback_data=f"exp:save:{pid}")],
        [InlineKeyboardButton("❌ Отмена",     callback_data=f"exp:cancel:{pid}")],
    ])


# ---------- views ----------

async def project_view(pid):
    async with pool.acquire() as c:
        p = await c.fetchrow("SELECT * FROM projects WHERE id=$1", pid)
        if not p:
            return None, None
        t = await c.fetchrow(
            "SELECT COALESCE(SUM(qty*price),0) AS total FROM expenses WHERE project_id=$1",
            pid,
        )
    total = Decimal(str(t["total"]))
    budget = Decimal(str(p["budget"]))
    left = budget - total
    status = "🟢" if left >= 0 else "🔴"
    text = (
        f"📁 <b>{esc(p['name'])}</b>\n\n"
        f"💰 Бюджет: {money(budget)}\n"
        f"💸 Потрачено: {money(total)}\n"
        f"{status} Осталось: {money(left)}"
    )
    return text, kb_project(pid)


async def send_project_msg(target, pid, edit=False):
    text, kb = await project_view(pid)
    if text is None:
        msg = "Проект не найден."
        return await (target.edit_message_text(msg) if edit else target.reply_text(msg))
    if edit:
        return await target.edit_message_text(text, parse_mode=ParseMode.HTML, reply_markup=kb)
    return await target.reply_text(text, parse_mode=ParseMode.HTML, reply_markup=kb)


async def projects_list_view():
    async with pool.acquire() as c:
        rows = await c.fetch(
            """
            SELECT p.id, p.name, p.budget, p.created_at,
                   COALESCE(SUM(e.qty*e.price),0) AS spent,
                   COUNT(e.id) AS cnt
            FROM projects p
            LEFT JOIN expenses e ON e.project_id = p.id
            GROUP BY p.id
            ORDER BY p.created_at DESC
            """
        )
    if not rows:
        text = "📁 <b>Общие проекты</b>\n\nПока нет ни одного проекта."
        kb = InlineKeyboardMarkup([
            [InlineKeyboardButton("➕ Новый проект", callback_data="project:new")],
            [InlineKeyboardButton("🏠 Меню",         callback_data="menu")],
        ])
        return text, kb

    lines = ["📁 <b>Общие проекты</b>\n"]
    for r in rows:
        total = Decimal(str(r["spent"]))
        left = Decimal(str(r["budget"])) - total
        mark = "🟢" if left >= 0 else "🔴"
        lines.append(
            f"• <b>{esc(r['name'])}</b> — {money(total)} / {money(r['budget'])} {mark}\n"
            f"  расходов: {r['cnt']}"
        )

    kb_rows = [
        [InlineKeyboardButton(f"📁 {r['name'][:40]}", callback_data=f"project:view:{r['id']}")]
        for r in rows
    ]
    kb_rows.append([InlineKeyboardButton("➕ Новый проект", callback_data="project:new")])
    kb_rows.append([InlineKeyboardButton("🏠 Меню",         callback_data="menu")])
    return "\n".join(lines), InlineKeyboardMarkup(kb_rows)


def _exp_summary(exp):
    qv = Decimal(str(exp.get("qty", "1")))
    pv = Decimal(str(exp.get("price", "0")))
    s = qv * pv
    lines = [
        "<b>Проверьте расход:</b>",
        f"• Категория: {esc(exp.get('category', '—'))}",
        f"• Название: {esc(exp.get('name', '—'))}",
        f"• Количество: {fmt_qty(qv)}",
        f"• Цена: {money(pv)}",
        f"• Сумма: {money(s)}",
    ]
    if exp.get("comment"):
        lines.append(f"• Комментарий: {esc(exp['comment'])}")
    return "\n".join(lines)


# ---------- Excel / PDF ----------

def _safe_name(name: str) -> str:
    s = re.sub(r"[^\w\-]+", "_", name or "").strip("_")
    return (s[:40] or "smeta")


async def make_excel(update, pid):
    q = update.callback_query
    async with pool.acquire() as c:
        p = await c.fetchrow("SELECT * FROM projects WHERE id=$1", pid)
        rows = await c.fetch("SELECT * FROM expenses WHERE project_id=$1 ORDER BY id", pid)
    if not p:
        return await q.message.reply_text("Проект не найден.")

    total = sum((Decimal(str(r["qty"])) * Decimal(str(r["price"])) for r in rows), Decimal(0))
    budget = Decimal(str(p["budget"]))
    left = budget - total

    wb = Workbook()
    ws = wb.active
    ws.title = "Смета"

    ws["A1"] = "СМЕТА"
    ws["A1"].font = Font(bold=True, size=16)
    ws.merge_cells("A1:H1")

    ws["A2"] = "Проект";     ws["B2"] = p["name"]
    ws["A3"] = "Бюджет";     ws["B3"] = float(budget)
    ws["A4"] = "Потрачено";  ws["B4"] = float(total)
    ws["A5"] = "Осталось";   ws["B5"] = float(left)

    headers = [
        "Категория", "Наименование", "Количество", "Цена", "Сумма",
        "Комментарий", "Кто добавил", "Дата",
    ]
    ws.append([])
    ws.append(headers)
    header_row = ws.max_row
    fill = PatternFill("solid", fgColor="DDDDDD")
    bold = Font(bold=True)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row, column=col)
        cell.font = bold
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")

    for r in rows:
        qv = Decimal(str(r["qty"]))
        pv = Decimal(str(r["price"]))
        s = qv * pv
        dt = r["created_at"].strftime("%d.%m.%Y %H:%M") if r["created_at"] else ""
        author = ("@" + r["author_username"]) if r["author_username"] else (
            f"id{r['author_id']}" if r["author_id"] else ""
        )
        ws.append([
            r["category"], r["name"], float(qv), float(pv), float(s),
            r["comment"] or "", author, dt,
        ])

    total_row = ws.max_row + 1
    ws.cell(row=total_row, column=4, value="ИТОГО").font = bold
    ws.cell(row=total_row, column=5, value=float(total)).font = bold

    for i, w in enumerate([18, 28, 12, 12, 14, 30, 18, 18], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    await q.message.reply_document(buf, filename=f"smeta_{_safe_name(p['name'])}.xlsx")


async def make_pdf(update, pid):
    q = update.callback_query
    async with pool.acquire() as c:
        p = await c.fetchrow("SELECT * FROM projects WHERE id=$1", pid)
        rows = await c.fetch("SELECT * FROM expenses WHERE project_id=$1 ORDER BY id", pid)
    if not p:
        return await q.message.reply_text("Проект не найден.")

    total = sum((Decimal(str(r["qty"])) * Decimal(str(r["price"])) for r in rows), Decimal(0))
    budget = Decimal(str(p["budget"]))
    left = budget - total

    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out, pagesize=A4,
        leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30,
    )
    styles = getSampleStyleSheet()
    body = styles["Normal"]

    story = [
        Paragraph("СМЕТА", styles["Title"]),
        Paragraph(esc(p["name"]), styles["Heading2"]),
        Spacer(1, 6),
        Paragraph(f"Бюджет: {money(budget)}", body),
        Paragraph(f"Потрачено: {money(total)}", body),
        Paragraph(f"Осталось: {money(left)}", body),
        Spacer(1, 12),
    ]

    data = [["Категория", "Наименование", "Кол-во", "Цена", "Сумма"]]
    for r in rows:
        qv = Decimal(str(r["qty"]))
        pv = Decimal(str(r["price"]))
        s = qv * pv
        data.append([r["category"], r["name"], fmt_qty(qv), money(pv), money(s)])
    data.append(["", "", "", "ИТОГО", money(total)])

    table = Table(data, repeatRows=1, colWidths=[95, 170, 55, 75, 85])
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), .4, colors.grey),
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(table)

    doc.build(story)
    out.seek(0)
    await q.message.reply_document(out, filename=f"smeta_{_safe_name(p['name'])}.pdf")


# ---------- commands ----------

async def cmd_start(update, ctx):
    u = update.effective_user
    await upsert_user(u)
    text = (
        f"👋 Привет, {esc(u.first_name or 'друг')}!\n\n"
        "Это бот для <b>командной</b> работы со сметами. "
        "Все проекты — общие: любой пользователь бота видит их, "
        "может открыть, дополнить расходом и скачать Excel/PDF.\n\n"
        "Выберите действие:"
    )
    await update.message.reply_text(text, parse_mode=ParseMode.HTML, reply_markup=kb_main())


async def cmd_menu(update, ctx):
    await upsert_user(update.effective_user)
    await update.message.reply_text("Меню:", reply_markup=kb_main())


async def cmd_projects(update, ctx):
    await upsert_user(update.effective_user)
    text, kb = await projects_list_view()
    await update.message.reply_text(text, parse_mode=ParseMode.HTML, reply_markup=kb)


async def cmd_cancel(update, ctx):
    for k in ("state", "exp", "new_project", "budget_pid"):
        ctx.user_data.pop(k, None)
    await update.message.reply_text("Отменено.", reply_markup=kb_main())


# ---------- text handler ----------

async def on_text(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    u = update.effective_user
    await upsert_user(u)
    state = ctx.user_data.get("state")
    text = (update.message.text or "").strip()

    if not state:
        # Если пользователь уже внутри проекта — пробуем быстрый ввод
        pid = ctx.user_data.get("current_project")
        if pid:
            parsed = parse_quick(text)
            if parsed:
                name, qty, price = parsed
                ctx.user_data["exp"] = {
                    "project_id": pid,
                    "name": name, "qty": str(qty), "price": str(price),
                }
                ctx.user_data["state"] = "exp:category"
                total = qty * price
                return await update.message.reply_text(
                    "Разобрал:\n"
                    f"• Название: {esc(name)}\n"
                    f"• Количество: {fmt_qty(qty)}\n"
                    f"• Цена: {money(price)}\n"
                    f"• Сумма: {money(total)}\n\n"
                    "Выберите категорию:",
                    parse_mode=ParseMode.HTML,
                    reply_markup=kb_categories(pid),
                )
        return await update.message.reply_text(
            "Используйте /start для меню.",
            reply_markup=kb_main(),
        )

    if text.lower() in ("/cancel", "отмена"):
        for k in ("state", "exp", "new_project", "budget_pid"):
            ctx.user_data.pop(k, None)
        return await update.message.reply_text("Отменено.", reply_markup=kb_main())

    # --- новый проект: название ---
    if state == "new_project:name":
        if not text:
            return
        ctx.user_data["new_project"] = {"name": text}
        ctx.user_data["state"] = "new_project:budget"
        return await update.message.reply_text(
            f"Проект: <b>{esc(text)}</b>\n\n"
            "Отправьте бюджет числом (₽). Например: 100000.\n"
            "Если бюджет неизвестен — отправьте 0.",
            parse_mode=ParseMode.HTML,
        )

    # --- новый проект: бюджет ---
    if state == "new_project:budget":
        val = parse_money(text)
        if val is None:
            return await update.message.reply_text("Не понял. Отправьте число, например 100000.")
        np = ctx.user_data.get("new_project") or {}
        async with pool.acquire() as c:
            row = await c.fetchrow(
                "INSERT INTO projects(name, budget, creator_id) VALUES($1,$2,$3) RETURNING id",
                np.get("name", "Без названия"), val, u.id,
            )
            pid = row["id"]
            await c.execute(
                "INSERT INTO project_members(project_id, user_id, role) VALUES($1,$2,$3) "
                "ON CONFLICT DO NOTHING",
                pid, u.id, ROLE_OWNER,
            )
        ctx.user_data.pop("state", None)
        ctx.user_data.pop("new_project", None)
        ctx.user_data["current_project"] = pid
        await update.message.reply_text(f"✅ Проект «{esc(np.get('name',''))}» создан.",
                                        parse_mode=ParseMode.HTML)
        return await send_project_msg(update.message, pid)

    # --- расход: название / быстрый ввод ---
    if state == "exp:name":
        exp = ctx.user_data.get("exp") or {}
        parsed = parse_quick(text)
        if parsed:
            name, qty, price = parsed
            exp.update({"name": name, "qty": str(qty), "price": str(price)})
            ctx.user_data["exp"] = exp
            ctx.user_data["state"] = "exp:category"
            return await update.message.reply_text(
                "Разобрал:\n"
                f"• Название: {esc(name)}\n"
                f"• Количество: {fmt_qty(qty)}\n"
                f"• Цена: {money(price)}\n"
                f"• Сумма: {money(qty*price)}\n\n"
                "Выберите категорию:",
                parse_mode=ParseMode.HTML,
                reply_markup=kb_categories(exp["project_id"]),
            )
        exp["name"] = text
        ctx.user_data["exp"] = exp
        ctx.user_data["state"] = "exp:qty"
        return await update.message.reply_text(
            f"Название: <b>{esc(text)}</b>\n\nСколько? (например 1, 3, 2.5)",
            parse_mode=ParseMode.HTML,
        )

    # --- расход: количество ---
    if state == "exp:qty":
        val = parse_number(text)
        if val is None:
            return await update.message.reply_text("Нужно число. Например: 1")
        exp = ctx.user_data["exp"]
        exp["qty"] = str(val)
        ctx.user_data["state"] = "exp:price"
        return await update.message.reply_text("Цена за единицу? (например 850)")

    # --- расход: цена ---
    if state == "exp:price":
        val = parse_money(text)
        if val is None:
            return await update.message.reply_text("Нужно число. Например: 850")
        exp = ctx.user_data["exp"]
        exp["price"] = str(val)
        ctx.user_data["state"] = "exp:category"
        total = Decimal(exp["qty"]) * Decimal(exp["price"])
        return await update.message.reply_text(
            f"Итого: <b>{money(total)}</b>\n\nВыберите категорию:",
            parse_mode=ParseMode.HTML,
            reply_markup=kb_categories(exp["project_id"]),
        )

    # --- расход: комментарий ---
    if state == "exp:comment":
        exp = ctx.user_data["exp"]
        exp["comment"] = text
        ctx.user_data["state"] = "exp:confirm"
        return await update.message.reply_text(
            _exp_summary(exp),
            parse_mode=ParseMode.HTML,
            reply_markup=kb_confirm(exp["project_id"]),
        )

    # --- редактирование бюджета ---
    if state == "project:budget":
        val = parse_money(text)
        if val is None:
            return await update.message.reply_text("Нужно число.")
        pid = ctx.user_data.get("budget_pid")
        async with pool.acquire() as c:
            await c.execute("UPDATE projects SET budget=$1 WHERE id=$2", val, pid)
        ctx.user_data.pop("state", None)
        ctx.user_data.pop("budget_pid", None)
        await update.message.reply_text("✅ Бюджет обновлён.")
        return await send_project_msg(update.message, pid)

    await update.message.reply_text("Не понял. /cancel — отмена.", reply_markup=kb_main())


# ---------- callback router ----------

async def callbacks(update: Update, ctx: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    await q.answer()
    data = q.data or ""
    u = update.effective_user
    await upsert_user(u)

    # список / меню
    if data == "menu":
        return await q.edit_message_text("Меню:", reply_markup=kb_main())
    if data == "projects":
        text, kb = await projects_list_view()
        return await q.edit_message_text(text, parse_mode=ParseMode.HTML, reply_markup=kb)

    # новый проект
    if data == "project:new":
        ctx.user_data["state"] = "new_project:name"
        return await q.edit_message_text(
            "➕ <b>Новый проект</b>\n\nОтправьте название проекта.\n\n/cancel — отмена",
            parse_mode=ParseMode.HTML,
            reply_markup=kb_back(),
        )

    # открыть проект
    if data.startswith("project:view:"):
        pid = int(data.split(":")[2])
        ctx.user_data["current_project"] = pid
        return await send_project_msg(q.message, pid, edit=True)

    # настройки
    if data.startswith("project:settings:"):
        pid = int(data.split(":")[2])
        async with pool.acquire() as c:
            p = await c.fetchrow("SELECT * FROM projects WHERE id=$1", pid)
            owner = await c.fetchrow(
                "SELECT username, first_name FROM users WHERE user_id=$1",
                p["creator_id"] if p else 0,
            ) if p else None
        if not p:
            return await q.edit_message_text("Проект не найден.")
        owner_name = (
            ("@" + owner["username"]) if owner and owner["username"]
            else (owner["first_name"] if owner and owner["first_name"] else f"id{p['creator_id']}")
        )
        created = p["created_at"].strftime("%d.%m.%Y") if p["created_at"] else ""
        text = (
            "⚙️ <b>Настройки проекта</b>\n\n"
            f"📁 {esc(p['name'])}\n"
            f"💰 Бюджет: {money(p['budget'])}\n"
            f"👑 Создатель: {esc(owner_name)}\n"
            f"📅 Создан: {created}\n\n"
            "Роли (владелец / участник / наблюдатель) уже заложены в БД "
            "и появятся в UI позже."
        )
        kb = InlineKeyboardMarkup([
            [InlineKeyboardButton("✏️ Изменить бюджет", callback_data=f"project:budget:{pid}")],
            [InlineKeyboardButton("◀️ Назад",           callback_data=f"project:view:{pid}")],
        ])
        return await q.edit_message_text(text, parse_mode=ParseMode.HTML, reply_markup=kb)

    if data.startswith("project:budget:"):
        pid = int(data.split(":")[2])
        ctx.user_data["state"] = "project:budget"
        ctx.user_data["budget_pid"] = pid
        return await q.edit_message_text(
            "Отправьте новый бюджет числом (например 150000).\n\n/cancel — отмена."
        )

    # удаление
    if data.startswith("project:delete:"):
        pid = int(data.split(":")[2])
        kb = InlineKeyboardMarkup([
            [InlineKeyboardButton("🗑 Да, удалить", callback_data=f"project:delete_yes:{pid}")],
            [InlineKeyboardButton("◀️ Отмена",      callback_data=f"project:view:{pid}")],
        ])
        return await q.edit_message_text(
            "Удалить проект и все его расходы? Действие необратимо.",
            reply_markup=kb,
        )
    if data.startswith("project:delete_yes:"):
        pid = int(data.split(":")[2])
        async with pool.acquire() as c:
            await c.execute("DELETE FROM projects WHERE id=$1", pid)
        ctx.user_data.pop("current_project", None)
        text, kb = await projects_list_view()
        return await q.edit_message_text(text, parse_mode=ParseMode.HTML, reply_markup=kb)

    # расход: старт
    if data.startswith("exp:add:"):
        pid = int(data.split(":")[2])
        ctx.user_data["state"] = "exp:name"
        ctx.user_data["exp"] = {"project_id": pid}
        ctx.user_data["current_project"] = pid
        return await q.edit_message_text(
            "➕ <b>Новый расход</b>\n\n"
            "Напишите одной строкой, например:\n"
            "• <code>такси 1200</code>\n"
            "• <code>3 банки краски по 850</code>\n"
            "• <code>свет 2 x 500</code>\n\n"
            "Или отправьте название — введу по шагам.\n\n"
            "/cancel — отмена",
            parse_mode=ParseMode.HTML,
            reply_markup=InlineKeyboardMarkup([
                [InlineKeyboardButton("❌ Отмена", callback_data=f"exp:cancel:{pid}")],
            ]),
        )

    # расход: список
    if data.startswith("exp:list:"):
        pid = int(data.split(":")[2])
        async with pool.acquire() as c:
            p = await c.fetchrow("SELECT * FROM projects WHERE id=$1", pid)
            rows = await c.fetch(
                "SELECT * FROM expenses WHERE project_id=$1 ORDER BY id DESC LIMIT 100",
                pid,
            )
        if not p:
            return await q.edit_message_text("Проект не найден.")
        if not rows:
            body = "Расходов пока нет."
        else:
            blocks = []
            for r in rows:
                qv = Decimal(str(r["qty"]))
                pv = Decimal(str(r["price"]))
                s = qv * pv
                author = (
                    ("@" + r["author_username"]) if r["author_username"]
                    else (f"id{r['author_id']}" if r["author_id"] else "—")
                )
                dt = r["created_at"].strftime("%d.%m.%Y %H:%M") if r["created_at"] else ""
                block = [
                    f"{esc(r['category'])}",
                    f"<b>{esc(r['name'])}</b>",
                    f"{fmt_qty(qv)} × {money(pv)} = {money(s)}",
                ]
                if r["comment"]:
                    block.append(f"💬 {esc(r['comment'])}")
                block.append(f"👤 {esc(author)} · {dt}")
                blocks.append("\n".join(block))
            body = "\n\n".join(blocks)
        text = f"📋 <b>Расходы проекта «{esc(p['name'])}»</b>\n\n{body}"
        if len(text) > 4000:
            text = text[:3900] + "\n\n… (показаны не все)"
        kb = InlineKeyboardMarkup([
            [InlineKeyboardButton("➕ Добавить расход",    callback_data=f"exp:add:{pid}")],
            [InlineKeyboardButton("◀️ Назад к проекту",    callback_data=f"project:view:{pid}")],
        ])
        return await q.edit_message_text(text, parse_mode=ParseMode.HTML, reply_markup=kb)

    # расход: категория
    if data.startswith("cat:"):
        _, pid_s, idx_s = data.split(":")
        pid = int(pid_s)
        cat = CATEGORIES[int(idx_s)]
        exp = ctx.user_data.get("exp") or {"project_id": pid}
        exp["category"] = cat
        ctx.user_data["exp"] = exp
        ctx.user_data["state"] = "exp:comment"
        return await q.edit_message_text(
            f"Категория: {esc(cat)}\n\nКомментарий? Отправьте текст или нажмите «Пропустить».",
            parse_mode=ParseMode.HTML,
            reply_markup=kb_comment(pid),
        )

    # расход: пропустить комментарий
    if data.startswith("exp:skip_comment:"):
        pid = int(data.split(":")[2])
        exp = ctx.user_data.get("exp") or {"project_id": pid}
        exp.setdefault("category", "📦 Другое")
        exp["comment"] = None
        ctx.user_data["exp"] = exp
        ctx.user_data["state"] = "exp:confirm"
        return await q.edit_message_text(
            _exp_summary(exp),
            parse_mode=ParseMode.HTML,
            reply_markup=kb_confirm(pid),
        )

    # расход: сохранить
    if data.startswith("exp:save:"):
        pid = int(data.split(":")[2])
        exp = ctx.user_data.get("exp") or {}
        name = exp.get("name")
        if not name:
            return await q.edit_message_text("Что-то пошло не так. Начните заново.")
        qty = Decimal(str(exp.get("qty", "1")))
        price = Decimal(str(exp.get("price", "0")))
        cat = exp.get("category", "📦 Другое")
        comment = exp.get("comment")
        async with pool.acquire() as c:
            await c.execute(
                """
                INSERT INTO expenses(project_id, category, name, qty, price,
                                     comment, author_id, author_username)
                VALUES($1,$2,$3,$4,$5,$6,$7,$8)
                """,
                pid, cat, name, qty, price, comment, u.id, u.username,
            )
            await c.execute(
                "INSERT INTO project_members(project_id, user_id, role) VALUES($1,$2,$3) "
                "ON CONFLICT DO NOTHING",
                pid, u.id, ROLE_MEMBER,
            )
        ctx.user_data.pop("state", None)
        ctx.user_data.pop("exp", None)
        ctx.user_data["current_project"] = pid
        text, kb = await project_view(pid)
        return await q.edit_message_text(
            "✅ Расход добавлен.\n\n" + text,
            parse_mode=ParseMode.HTML,
            reply_markup=kb,
        )

    # расход: отмена
    if data.startswith("exp:cancel:"):
        pid = int(data.split(":")[2])
        ctx.user_data.pop("state", None)
        ctx.user_data.pop("exp", None)
        text, kb = await project_view(pid)
        return await q.edit_message_text(
            "Отменено.\n\n" + text,
            parse_mode=ParseMode.HTML,
            reply_markup=kb,
        )

    # экспорт
    if data.startswith("excel:"):
        return await make_excel(update, int(data.split(":")[2]))
    if data.startswith("pdf:"):
        return await make_pdf(update, int(data.split(":")[2]))

    log.warning("Unhandled callback: %s", data)


# ---------- entrypoint ----------

async def post_init(app):
    await init_db()
    me = await app.bot.get_me()
    log.info("Bot @%s started", me.username)


def main():
    app = (
        Application.builder()
        .token(BOT_TOKEN)
        .post_init(post_init)
        .build()
    )
    app.add_handler(CommandHandler("start", cmd_start))
    app.add_handler(CommandHandler("menu", cmd_menu))
    app.add_handler(CommandHandler("projects", cmd_projects))
    app.add_handler(CommandHandler("cancel", cmd_cancel))
    app.add_handler(CallbackQueryHandler(callbacks))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, on_text))
    app.run_polling()


if __name__ == "__main__":
    main()
